"""
Generate a sample Excel file `modele_emploi_du_temps.xlsx` with curriculum and timetable sheets.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

def generate_sample_excel():
    wb = openpyxl.Workbook()

    # Remove default sheet
    default_sheet = wb.active

    # Style definitions
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # 1. Sheet Programs
    ws_programs = wb.create_sheet(title="Programs")
    ws_programs.append(["Program_Name", "Program_Code", "Description"])
    ws_programs.append(["Licence Informatique & IA", "INFO-IA", "Licence Informatique et Intelligence Artificielle"])
    ws_programs.append(["Master Data Science", "DS-MAST", "Master Science des Donnees"])

    # 2. Sheet ClassSchedules (THE TIMETABLE SHEET)
    ws_schedules = wb.create_sheet(title="ClassSchedules")
    headers = [
        "Program_Name",
        "Track_Name",
        "Semester_Number",
        "Course_Code",
        "Course_Name",
        "Day_Of_Week",
        "Start_Time",
        "End_Time",
        "Session_Type",
        "Room_Location",
        "Is_Mandatory"
    ]
    ws_schedules.append(headers)

    sample_schedules = [
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF101", "Algorithmique & Structures de Donnees", "Monday", "08:30", "10:30", "CM", "Amphi Alan Turing", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF101-TD", "Algorithmique (Travaux Diriges)", "Monday", "11:00", "13:00", "TD", "Salle TD 204", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF102", "Architecture des Ordinateurs", "Tuesday", "09:00", "12:00", "CM", "Amphi Neumann", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF103", "Bases de Donnees Relationnelles", "Wednesday", "08:30", "11:30", "CM", "Amphi Turing", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF103-TP", "Bases de Donnees (TP SQL)", "Wednesday", "14:00", "17:00", "TP", "Labo Info 3", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF104", "Developpement Web Fullstack", "Thursday", "10:00", "12:30", "CM", "Amphi Lovelace", "TRUE"],
        ["Licence Informatique & IA", "Software Engineering Track", 1, "INF104-TP", "Projet Web & Reseau", "Friday", "14:00", "17:00", "TP", "Labo Info 1", "TRUE"],
    ]

    for row in sample_schedules:
        ws_schedules.append(row)

    wb.remove(default_sheet)

    # Apply styling
    for sheet in wb.worksheets:
        for col_idx, col in enumerate(sheet.columns, 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.row_dimensions[1].height = 28

            # Column auto width
            max_len = max(len(str(c.value or '')) for c in col)
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(max_len + 4, 15)

            for row_idx in range(2, sheet.max_row + 1):
                c = sheet.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(vertical="center")

    output_path = Path("uploads") / "modele_emploi_du_temps.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Fichier modele cree avec succes : {output_path.absolute()}")

if __name__ == "__main__":
    generate_sample_excel()
